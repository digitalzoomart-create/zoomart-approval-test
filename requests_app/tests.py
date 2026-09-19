"""
Automated tests for the critical business rules called out in the
project spec (section 19): self-approval prevention, department
isolation, approval sequencing, rejection stopping the workflow,
amount-based routing, and audit logging.
"""

from django.contrib.auth.models import Group
from django.test import Client, TestCase

from accounts.models import User
from core.models import ApprovalStepRule, ApprovalWorkflowRule, Department, RequestCategory
from core.permissions import can_view_request
from . import services
from .models import Request, RequestApproval


class BaseWorkflowTestCase(TestCase):
    def setUp(self):
        for name in ["Employee", "Manager", "Finance", "Administrator", "Senior Management", "Procurement Manager"]:
            Group.objects.get_or_create(name=name)

        self.dept_a = Department.objects.create(name="Marketing")
        self.dept_b = Department.objects.create(name="IT")
        self.category = RequestCategory.objects.create(name="Office Supplies")

        self.manager_a = self._make_user("manager_a", ["Manager"], self.dept_a)
        self.manager_b = self._make_user("manager_b", ["Manager"], self.dept_b)
        self.finance_user = self._make_user("finance_user", ["Finance"])
        self.director = self._make_user("director", ["Senior Management"])
        self.procurement_manager = self._make_user("procurement_manager", ["Procurement Manager"])
        self.employee_a = self._make_user("employee_a", ["Employee"], self.dept_a)
        self.employee_a2 = self._make_user("employee_a2", ["Employee"], self.dept_a)
        self.employee_b = self._make_user("employee_b", ["Employee"], self.dept_b)

        self.dept_a.manager = self.manager_a
        self.dept_a.save()
        self.dept_b.manager = self.manager_b
        self.dept_b.save()

        # Small purchases: department director only.
        low = ApprovalWorkflowRule.objects.create(name="Low", min_amount=0, max_amount=500)
        ApprovalStepRule.objects.create(workflow_rule=low, order=1, approver_role=ApprovalStepRule.ROLE_DEPARTMENT_MANAGER)

        # Everything above: department director -> company director. Once the
        # director approves, the request is fully APPROVED — Finance executes
        # it (purchase/paid/completed) but does not re-approve it.
        standard = ApprovalWorkflowRule.objects.create(name="Standard", min_amount=500.01, max_amount=None)
        ApprovalStepRule.objects.create(workflow_rule=standard, order=1, approver_role=ApprovalStepRule.ROLE_DEPARTMENT_MANAGER)
        ApprovalStepRule.objects.create(workflow_rule=standard, order=2, approver_role=ApprovalStepRule.ROLE_SENIOR_MANAGER)

    def _make_user(self, username, roles, department=None):
        u = User.objects.create_user(username=username, password="pass12345", department=department)
        u.groups.set(Group.objects.filter(name__in=roles))
        return u

    def _make_request(self, requester, department, amount, title="Test request"):
        return Request.objects.create(
            requester=requester, department=department, category=self.category,
            title=title, description="desc", business_justification="because", estimated_cost=amount,
        )


class ApprovalRoutingTests(BaseWorkflowTestCase):
    def test_low_amount_only_needs_manager(self):
        req = self._make_request(self.employee_a, self.dept_a, 300)
        services.submit_request(req, self.employee_a)
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_PENDING_MANAGER)
        self.assertEqual(req.approval_steps.count(), 1)

        services.approve(req, self.manager_a, "ok")
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_APPROVED, "single-step chain should be fully approved")

    def test_mid_amount_needs_manager_then_director_then_becomes_approved(self):
        req = self._make_request(self.employee_a, self.dept_a, 1500)
        services.submit_request(req, self.employee_a)
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_PENDING_MANAGER)

        services.approve(req, self.manager_a, "ok")
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_PENDING_SENIOR, "must advance to the company director next, not skip it")

        services.approve(req, self.director, "ok")
        req.refresh_from_db()
        self.assertEqual(
            req.status, Request.STATUS_APPROVED,
            "the company director's approval is final — Finance does not re-approve it",
        )

    def test_high_amount_needs_both_approval_steps(self):
        req = self._make_request(self.employee_a, self.dept_a, 25000)
        services.submit_request(req, self.employee_a)
        services.approve(req, self.manager_a, "ok")
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_PENDING_SENIOR)
        services.approve(req, self.director, "ok")
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_APPROVED)

    def test_finance_never_has_a_pending_approval_step(self):
        req = self._make_request(self.employee_a, self.dept_a, 25000)
        services.submit_request(req, self.employee_a)
        services.approve(req, self.manager_a, "ok")
        services.approve(req, self.director, "ok")
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_APPROVED)
        self.assertFalse(
            req.approval_steps.filter(approver_role=ApprovalStepRule.ROLE_FINANCE).exists(),
            "the default workflow no longer includes a Finance approval step at all",
        )

    def test_finance_can_execute_but_not_approve_once_request_is_approved(self):
        req = self._make_request(self.employee_a, self.dept_a, 25000)
        services.submit_request(req, self.employee_a)
        services.approve(req, self.manager_a, "ok")
        services.approve(req, self.director, "ok")
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_APPROVED)

        # Finance has nothing to "approve" — the request is already approved —
        # but they can still execute the purchase lifecycle.
        with self.assertRaises(services.ApprovalError):
            services.approve(req, self.finance_user, "trying to approve anyway")

        services.mark_purchase_in_progress(req, self.finance_user)
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_PURCHASE_IN_PROGRESS)


class SelfSubmissionSkipTests(BaseWorkflowTestCase):
    def test_department_director_submitting_own_request_skips_that_step(self):
        req = self._make_request(self.manager_a, self.dept_a, 1500)
        services.submit_request(req, self.manager_a)
        req.refresh_from_db()
        self.assertEqual(
            req.status, Request.STATUS_PENDING_SENIOR,
            "the department-director step must be bypassed automatically and go straight to the company director",
        )
        dept_step = req.approval_steps.get(approver_role=ApprovalStepRule.ROLE_DEPARTMENT_MANAGER)
        self.assertEqual(dept_step.decision, RequestApproval.DECISION_SKIPPED)

        services.approve(req, self.director, "ok")
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_APPROVED)

    def test_department_director_submitting_own_small_request_auto_approves(self):
        # Only step in this tier is the department director's own — with
        # that step skipped, there is nothing left to approve.
        req = self._make_request(self.manager_a, self.dept_a, 300)
        services.submit_request(req, self.manager_a)
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_APPROVED)

    def test_company_director_submitting_own_request_with_no_department_auto_approves(self):
        # The company director is the highest approval rung — their own
        # request needs nobody's sign-off (and may have no department at
        # all), so it should go straight to APPROVED, ready for Finance.
        req = self._make_request(self.director, None, 1500)
        services.submit_request(req, self.director)
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_APPROVED)
        for step in req.approval_steps.all():
            self.assertEqual(step.decision, RequestApproval.DECISION_SKIPPED)

    def test_company_director_submitting_own_request_with_a_department_still_auto_approves(self):
        # Even if a department happens to be set, the director's own request
        # still skips every step — nobody outranks them to approve it.
        req = self._make_request(self.director, self.dept_a, 1500)
        services.submit_request(req, self.director)
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_APPROVED)


class ProcurementManagerOversightTests(BaseWorkflowTestCase):
    def test_procurement_manager_can_view_any_request(self):
        req = self._make_request(self.employee_b, self.dept_b, 300)
        self.assertTrue(can_view_request(self.procurement_manager, req))

    def test_procurement_manager_cannot_act_as_approver(self):
        from core.permissions import can_act_as_approver

        req = self._make_request(self.employee_a, self.dept_a, 1500)
        services.submit_request(req, self.employee_a)
        self.assertFalse(can_act_as_approver(self.procurement_manager, req))


class NotificationTests(BaseWorkflowTestCase):
    def test_department_director_is_emailed_when_their_approval_is_needed(self):
        from django.core import mail

        self.manager_a.email = "manager_a@zoomart.ge"
        self.manager_a.save()
        req = self._make_request(self.employee_a, self.dept_a, 1500)
        mail.outbox = []
        services.submit_request(req, self.employee_a)
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn(self.manager_a.email, mail.outbox[0].to)

    def test_company_director_is_emailed_only_once_it_is_their_turn(self):
        from django.core import mail

        self.manager_a.email = "manager_a@zoomart.ge"
        self.manager_a.save()
        self.director.email = "director@zoomart.ge"
        self.director.save()
        req = self._make_request(self.employee_a, self.dept_a, 1500)
        services.submit_request(req, self.employee_a)
        mail.outbox = []
        services.approve(req, self.manager_a, "ok")
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn(self.director.email, mail.outbox[0].to)

    def test_finance_is_not_emailed_before_the_request_is_fully_approved(self):
        from django.core import mail

        self.manager_a.email = "manager_a@zoomart.ge"
        self.manager_a.save()
        self.finance_user.email = "finance@zoomart.ge"
        self.finance_user.save()
        req = self._make_request(self.employee_a, self.dept_a, 1500)
        mail.outbox = []
        services.submit_request(req, self.employee_a)
        for msg in mail.outbox:
            self.assertNotIn(self.finance_user.email, msg.to)

    def test_finance_is_emailed_once_request_is_fully_approved(self):
        from django.core import mail

        self.manager_a.email = "manager_a@zoomart.ge"
        self.manager_a.save()
        self.director.email = "director@zoomart.ge"
        self.director.save()
        self.finance_user.email = "finance@zoomart.ge"
        self.finance_user.save()
        req = self._make_request(self.employee_a, self.dept_a, 1500)
        services.submit_request(req, self.employee_a)
        services.approve(req, self.manager_a, "ok")
        mail.outbox = []
        services.approve(req, self.director, "ok")
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_APPROVED)
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn(self.finance_user.email, mail.outbox[0].to)

    def test_finance_is_emailed_when_company_directors_own_request_auto_approves(self):
        from django.core import mail

        self.finance_user.email = "finance@zoomart.ge"
        self.finance_user.save()
        req = self._make_request(self.director, None, 1500)
        mail.outbox = []
        services.submit_request(req, self.director)
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn(self.finance_user.email, mail.outbox[0].to)


class SelfApprovalTests(BaseWorkflowTestCase):
    def test_requester_cannot_approve_own_request(self):
        req = self._make_request(self.employee_a, self.dept_a, 300)
        services.submit_request(req, self.employee_a)
        with self.assertRaises(services.ApprovalError):
            services.approve(req, self.employee_a, "sneaky self-approval")
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_PENDING_MANAGER, "status must not change")

    def test_requester_cannot_reject_own_request(self):
        req = self._make_request(self.employee_a, self.dept_a, 300)
        services.submit_request(req, self.employee_a)
        with self.assertRaises(services.ApprovalError):
            services.reject(req, self.employee_a, "trying to reject my own request")

    def test_department_manager_of_a_different_department_cannot_approve(self):
        req = self._make_request(self.employee_a, self.dept_a, 300)
        services.submit_request(req, self.employee_a)
        with self.assertRaises(services.ApprovalError):
            services.approve(req, self.manager_b, "wrong department manager")


class RejectionAndInfoRequestTests(BaseWorkflowTestCase):
    def test_rejection_requires_comment(self):
        req = self._make_request(self.employee_a, self.dept_a, 300)
        services.submit_request(req, self.employee_a)
        with self.assertRaises(services.ApprovalError):
            services.reject(req, self.manager_a, "")

    def test_rejection_stops_the_workflow_and_skips_remaining_steps(self):
        req = self._make_request(self.employee_a, self.dept_a, 1500)
        services.submit_request(req, self.employee_a)
        services.reject(req, self.manager_a, "Not needed right now.")
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_REJECTED)
        remaining = req.approval_steps.filter(decision=RequestApproval.DECISION_PENDING)
        self.assertEqual(remaining.count(), 0, "no step should still be pending after a rejection")
        senior_step = req.approval_steps.get(approver_role=ApprovalStepRule.ROLE_SENIOR_MANAGER)
        self.assertEqual(senior_step.decision, RequestApproval.DECISION_SKIPPED)

    def test_more_info_then_resubmit_returns_to_same_step(self):
        req = self._make_request(self.employee_a, self.dept_a, 300)
        services.submit_request(req, self.employee_a)
        services.request_more_info(req, self.manager_a, "Please attach a quotation.")
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_MORE_INFO)

        services.submit_request(req, self.employee_a)
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_PENDING_MANAGER)
        self.assertEqual(req.approval_steps.count(), 1, "should not create a duplicate step chain")

        services.approve(req, self.manager_a, "ok now")
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_APPROVED)


class VisibilityAndPermissionTests(BaseWorkflowTestCase):
    def test_employee_cannot_see_another_departments_request(self):
        req = self._make_request(self.employee_b, self.dept_b, 300)
        self.assertFalse(can_view_request(self.employee_a, req))

    def test_employee_can_see_own_request(self):
        req = self._make_request(self.employee_a, self.dept_a, 300)
        self.assertTrue(can_view_request(self.employee_a, req))

    def test_manager_can_see_their_own_departments_requests(self):
        req = self._make_request(self.employee_a, self.dept_a, 300)
        self.assertTrue(can_view_request(self.manager_a, req))

    def test_manager_cannot_see_another_departments_requests(self):
        req = self._make_request(self.employee_b, self.dept_b, 300)
        self.assertFalse(can_view_request(self.manager_a, req))

    def test_confidential_request_hidden_from_department_manager_without_a_role_in_it(self):
        req = self._make_request(self.employee_a, self.dept_a, 300)
        req.is_confidential = True
        req.save()
        self.assertFalse(can_view_request(self.manager_a, req), "confidential requests are opt-in visible only")

    def test_administrator_sees_everything(self):
        admin = self._make_user("admin_user", ["Administrator"])
        admin.is_superuser = True
        admin.save()
        req = self._make_request(self.employee_b, self.dept_b, 300)
        self.assertTrue(can_view_request(admin, req))

    def test_web_view_returns_404_for_unauthorized_department(self):
        req = self._make_request(self.employee_b, self.dept_b, 300)
        client = Client()
        client.force_login(self.employee_a)
        resp = client.get(f"/requests/{req.pk}/")
        self.assertEqual(resp.status_code, 404)


class AuditLogTests(BaseWorkflowTestCase):
    def test_creating_and_approving_a_request_writes_audit_log_entries(self):
        from auditlog.models import LogEntry

        before = LogEntry.objects.count()
        req = self._make_request(self.employee_a, self.dept_a, 300)
        services.submit_request(req, self.employee_a)
        services.approve(req, self.manager_a, "ok")
        after = LogEntry.objects.count()
        self.assertGreater(after, before, "audit log entries should be recorded for create/update actions")


class PendingForUserTests(BaseWorkflowTestCase):
    """Regression test: a future (not-yet-reached) step must not show up as
    'waiting for my approval' just because it defaults to a pending-like state."""

    def test_future_step_does_not_appear_before_the_current_one_is_decided(self):
        # A Finance approval step is no longer part of the default seeded
        # workflow (Finance executes but doesn't approve — see
        # ApprovalRoutingTests), but ApprovalStepRule still supports one for
        # admins who want it, so this regression case builds its own 3-step
        # rule locally to keep covering that underlying engine behavior.
        custom_category = RequestCategory.objects.create(name="Custom rule category")
        rule = ApprovalWorkflowRule.objects.create(
            name="Manager-Finance-Senior", category=custom_category, min_amount=0, max_amount=None, priority=100,
        )
        ApprovalStepRule.objects.create(workflow_rule=rule, order=1, approver_role=ApprovalStepRule.ROLE_DEPARTMENT_MANAGER)
        ApprovalStepRule.objects.create(workflow_rule=rule, order=2, approver_role=ApprovalStepRule.ROLE_FINANCE)
        ApprovalStepRule.objects.create(workflow_rule=rule, order=3, approver_role=ApprovalStepRule.ROLE_SENIOR_MANAGER)

        req = Request.objects.create(
            requester=self.employee_a, department=self.dept_a, category=custom_category,
            title="Custom rule request", description="desc", business_justification="because", estimated_cost=5000,
        )
        services.submit_request(req, self.employee_a)
        services.approve(req, self.manager_a, "ok")
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_PENDING_FINANCE)

        director_queue = RequestApproval.objects.pending_for_user(self.director)
        self.assertNotIn(
            req.pk,
            director_queue.values_list("request_id", flat=True),
            "the company-director step exists in the DB already but is not active yet, "
            "so it must not appear in their queue while Finance's step is still pending",
        )

        finance_queue = RequestApproval.objects.pending_for_user(self.finance_user)
        self.assertIn(req.pk, finance_queue.values_list("request_id", flat=True))


class DepartmentDirectorRequestListTests(BaseWorkflowTestCase):
    """Regression test for the bug where a department director approving (or
    rejecting) a request made it disappear from their view entirely — it was
    only ever shown in their 'pending approvals' queue, which drops it the
    moment they decide on it. scope=department must keep it visible, with
    its current status, going forward."""

    def test_manager_still_sees_request_in_department_scope_after_approving(self):
        req = self._make_request(self.employee_a, self.dept_a, 300)
        services.submit_request(req, self.employee_a)

        client = Client()
        client.force_login(self.manager_a)

        # Before deciding: shows up as pending in scope=approvals.
        resp = client.get("/requests/?scope=approvals")
        self.assertContains(resp, req.request_number)

        services.approve(req, self.manager_a, "ok")
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_APPROVED)

        # It has left the pending-approvals queue now that it's decided...
        resp = client.get("/requests/?scope=approvals")
        self.assertNotContains(resp, req.request_number)

        # ...but must still be findable, with its status, in the
        # department director's own department view.
        resp = client.get("/requests/?scope=department")
        self.assertContains(resp, req.request_number)
        self.assertContains(resp, "დამტკიცებული")

    def test_manager_still_sees_request_in_department_scope_after_rejecting(self):
        req = self._make_request(self.employee_a, self.dept_a, 300)
        services.submit_request(req, self.employee_a)

        client = Client()
        client.force_login(self.manager_a)

        services.reject(req, self.manager_a, "Not approved.")
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_REJECTED)

        resp = client.get("/requests/?scope=department")
        self.assertContains(resp, req.request_number)

    def test_manager_cannot_see_another_departments_requests_in_department_scope(self):
        req = self._make_request(self.employee_b, self.dept_b, 300)
        services.submit_request(req, self.employee_b)
        services.approve(req, self.manager_b, "ok")

        client = Client()
        client.force_login(self.manager_a)
        resp = client.get("/requests/?scope=department")
        self.assertNotContains(resp, req.request_number)

    def test_non_manager_department_scope_falls_back_to_own_requests(self):
        req = self._make_request(self.employee_a, self.dept_a, 300)
        client = Client()
        client.force_login(self.employee_a)
        resp = client.get("/requests/?scope=department")
        # Employees aren't department directors, so this scope silently
        # falls back to the normal "my requests" view rather than erroring.
        self.assertContains(resp, req.request_number)


class RequestListExportTests(BaseWorkflowTestCase):
    """Excel export must be offered to department directors and Finance
    (and other elevated roles), must reflect the same scope/filters as the
    list the user is looking at, and must be refused to plain employees."""

    def test_manager_can_export_department_scope_as_xlsx(self):
        req = self._make_request(self.employee_a, self.dept_a, 300)
        services.submit_request(req, self.employee_a)
        services.approve(req, self.manager_a, "ok")

        client = Client()
        client.force_login(self.manager_a)
        resp = client.get("/requests/export/?scope=department")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        from io import BytesIO
        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(b"".join(resp.streaming_content)))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "მოთხოვნის №")
        self.assertIn(req.request_number, [row[0] for row in rows[1:]])

    def test_finance_user_can_export(self):
        client = Client()
        client.force_login(self.finance_user)
        resp = client.get("/requests/export/?scope=finance")
        self.assertEqual(resp.status_code, 200)

    def test_plain_employee_cannot_export(self):
        client = Client()
        client.force_login(self.employee_a)
        resp = client.get("/requests/export/")
        self.assertEqual(resp.status_code, 404)

    def test_export_button_shown_to_manager_and_finance_but_not_employee(self):
        client = Client()
        client.force_login(self.manager_a)
        resp = client.get("/requests/?scope=department")
        self.assertContains(resp, "ექსელში გადმოწერა")

        client.force_login(self.employee_a)
        resp = client.get("/requests/?scope=mine")
        self.assertNotContains(resp, "ექსელში გადმოწერა")


class RequestNumberAndDateFilterTests(BaseWorkflowTestCase):
    """Regression test for GIO's request: request numbers should encode the
    date they were created on, and the list should be filterable by date."""

    def test_request_number_encodes_todays_date(self):
        from django.utils import timezone

        req = self._make_request(self.employee_a, self.dept_a, 300)
        expected_prefix = "REQ-" + timezone.now().strftime("%Y%m%d") + "-"
        self.assertTrue(
            req.request_number.startswith(expected_prefix),
            f"expected {req.request_number} to start with {expected_prefix}",
        )

    def test_same_day_requests_get_sequential_numbers(self):
        req1 = self._make_request(self.employee_a, self.dept_a, 300, title="First")
        req2 = self._make_request(self.employee_a, self.dept_a, 300, title="Second")
        seq1 = int(req1.request_number.split("-")[-1])
        seq2 = int(req2.request_number.split("-")[-1])
        self.assertEqual(seq2, seq1 + 1)

    def test_list_can_be_filtered_by_date_range(self):
        import datetime

        req = self._make_request(self.employee_a, self.dept_a, 300)
        client = Client()
        client.force_login(self.employee_a)

        today = datetime.date.today()
        resp = client.get(f"/requests/?scope=mine&date_from={today.isoformat()}&date_to={today.isoformat()}")
        self.assertContains(resp, req.request_number)

        yesterday = today - datetime.timedelta(days=1)
        resp = client.get(f"/requests/?scope=mine&date_from={yesterday.isoformat()}&date_to={yesterday.isoformat()}")
        self.assertNotContains(resp, req.request_number)


class RaceConditionTests(BaseWorkflowTestCase):
    def test_cannot_approve_the_same_step_twice(self):
        req = self._make_request(self.employee_a, self.dept_a, 300)
        services.submit_request(req, self.employee_a)
        services.approve(req, self.manager_a, "first approval")
        req.refresh_from_db()
        self.assertEqual(req.status, Request.STATUS_APPROVED)
        # A second approval attempt must fail cleanly, not double-advance the workflow.
        with self.assertRaises(services.ApprovalError):
            services.approve(req, self.manager_a, "second approval attempt")


class RequestFormDepartmentFieldTests(BaseWorkflowTestCase):
    """The department field must be locked to the user's own department for
    employees/department directors, and optional (free choice) for the
    company director and other company-wide roles."""

    def test_department_field_locked_and_prefilled_for_employee(self):
        from .forms import RequestForm

        form = RequestForm(user=self.employee_a)
        self.assertTrue(form.fields["department"].disabled)
        self.assertEqual(form.initial["department"], self.dept_a.id)
        self.assertEqual(list(form.fields["department"].queryset), [self.dept_a])

    def test_department_field_not_required_for_company_director(self):
        from .forms import RequestForm

        form = RequestForm(user=self.director)
        self.assertFalse(form.fields["department"].disabled)
        self.assertFalse(form.fields["department"].required)

    def test_employee_cannot_submit_a_different_department_via_post(self):
        # Even if a manipulated request tried to smuggle a different
        # department id in, a disabled field ignores POST data entirely and
        # falls back to its initial value.
        from .forms import RequestForm

        data = {
            "title": "t", "description": "d", "department": self.dept_b.id,
            "category": self.category.id, "request_type": "PURCHASE", "priority": "MEDIUM",
            "estimated_cost": "10", "currency": "GEL", "quantity": 1,
            "business_justification": "because",
        }
        form = RequestForm(data, user=self.employee_a)
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["department"], self.dept_a)
